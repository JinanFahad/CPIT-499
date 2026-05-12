"""Generates a Word document containing the System Testing section
for the Muqaddim platform — organized by user role with scenario tables
modeled after the standard format (No, Scenario, Steps, Result)."""

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# ──────────────────────────────────────────────────────────────────────
# All system test scenarios, grouped by user role.
# Each scenario: (No, Scenario, Steps, Result)
# ──────────────────────────────────────────────────────────────────────
SCENARIOS = {
    "6.6.1 Idea Owner / End User System Test": [
        (
            "1",
            "Sign Up & Login",
            "Open Muqaddim > Click 'Sign Up' > Enter name, email, "
            "password > Click 'Register' > Verify email > Login with credentials",
            "Pass",
        ),
        (
            "2",
            "Generate Feasibility Study",
            "Login > Dashboard > Click 'New Feasibility Study' > "
            "Select business type (cafe) > Enter capital, rent, employees, "
            "average price, daily customers > Click 'Calculate' > "
            "System displays full feasibility report on screen",
            "Pass",
        ),
        (
            "3",
            "Download Feasibility Report as PDF",
            "From feasibility result page > Click 'Download PDF' > "
            "PDF file is generated and downloaded > Open file to verify content",
            "Pass",
        ),
        (
            "4",
            "Send Feasibility Report to Email",
            "From feasibility result page > Click 'Send to Email' > "
            "Enter email address > Click 'Send' > Email arrives "
            "with PDF attached",
            "Pass",
        ),
        (
            "5",
            "Generate Pitch Deck",
            "From a saved project > Click 'Generate Pitch Deck' > "
            "System generates PowerPoint file > Download and open in PowerPoint",
            "Pass",
        ),
        (
            "6",
            "Send Pitch Deck to Email",
            "From pitch deck page > Click 'Send to Email' > "
            "Enter email address > Click 'Send' > Email arrives with PPT attached",
            "Pass",
        ),
        (
            "7",
            "Market Analysis with Location",
            "From feasibility page > Click 'Pick Location on Map' > "
            "Select location in Riyadh > Confirm > System analyzes market "
            "and displays market score with nearby competitors",
            "Pass",
        ),
        (
            "8",
            "Chat with AI Consultant",
            "Login > Click 'AI Consultant' > Type business question > "
            "Send > AI responds with relevant guidance",
            "Pass",
        ),
        (
            "9",
            "Government Procedures Assistant",
            "Login > Click 'Government Procedures' > Ask about commercial "
            "registration > System responds with steps based on Saudi regulations",
            "Pass",
        ),
        (
            "10",
            "View Saved Projects",
            "Login > Click 'My Projects' > System displays list of "
            "all previously saved feasibility studies > Click on one to view details",
            "Pass",
        ),
        (
            "11",
            "Edit Existing Project",
            "From My Projects > Select a project > Click 'Edit' > "
            "Modify input values > Click 'Save' > System updates "
            "the project and recalculates results",
            "Pass",
        ),
        (
            "12",
            "Delete Project",
            "From My Projects > Select a project > Click 'Delete' > "
            "Confirm deletion > Project is removed from the list",
            "Pass",
        ),
        (
            "13",
            "Update Profile Information",
            "Login > Click 'Profile' > Edit name or password > "
            "Click 'Save' > Changes are saved successfully",
            "Pass",
        ),
        (
            "14",
            "Invalid Input Handling",
            "Open feasibility page > Enter capital = 0 or empty fields > "
            "Click 'Calculate' > System displays clear Arabic error messages",
            "Pass",
        ),
        (
            "15",
            "Logout",
            "Login > Click profile menu > Click 'Logout' > "
            "User is redirected to landing page > Cannot access protected pages",
            "Pass",
        ),
    ],
}


# ──────────────────────────────────────────────────────────────────────
# Word document construction helpers
# ──────────────────────────────────────────────────────────────────────
def set_cell_background(cell, hex_color: str) -> None:
    """Apply a background color to a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)


def add_paragraph(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.bold = bold


def add_scenario_table(doc: Document, rows: list) -> None:
    """Builds a 4-column table: No | Scenario | Steps | Result."""
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header row
    headers = ["No.", "Scenario", "Steps", "Result"]
    header_cells = table.rows[0].cells
    for idx, name in enumerate(headers):
        header_cells[idx].text = ""
        para = header_cells[idx].paragraphs[0]
        run = para.add_run(name)
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_background(header_cells[idx], "1F4E79")

    # Data rows
    for no, scenario, steps, result in rows:
        row_cells = table.add_row().cells
        row_cells[0].text = no
        row_cells[1].text = scenario
        row_cells[2].text = steps
        row_cells[3].text = result

        for cell in row_cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(10)

        # Color the Pass/Fail cell
        result_cell = row_cells[3]
        for para in result_cell.paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                run.bold = True
                if result.lower() == "pass":
                    run.font.color.rgb = RGBColor(0x00, 0x80, 0x00)
                else:
                    run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

    # Column widths
    widths = [Cm(1.2), Cm(4.0), Cm(9.5), Cm(1.8)]
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width


# ──────────────────────────────────────────────────────────────────────
# Build the document
# ──────────────────────────────────────────────────────────────────────
def build_document(output_path: str) -> None:
    doc = Document()

    # Section title
    add_heading(doc, "6.6 System Testing", level=1)

    # Intro paragraph
    intro = (
        "This section presents the system-level testing conducted for the "
        "Muqaddim platform. Each test evaluates a complete user workflow "
        "from end to end, validating expected behaviors, system responses, "
        "and overall workflow accuracy. The scenarios below cover all "
        "major features available to the end user."
    )
    add_paragraph(doc, intro)
    doc.add_paragraph()

    # Render each role's scenarios
    for section_title, rows in SCENARIOS.items():
        add_heading(doc, section_title, level=2)
        add_scenario_table(doc, rows)

        # Caption under the table
        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_run = caption.add_run(
            f"Table X: {section_title.split(' ', 1)[1]}"
        )
        caption_run.italic = True
        caption_run.font.size = Pt(10)
        doc.add_paragraph()

    doc.save(output_path)
    print(f"Document saved to: {output_path}")


if __name__ == "__main__":
    build_document("System_Testing_Muqaddim.docx")
